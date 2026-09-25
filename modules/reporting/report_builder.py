"""Moteur de publication Word du rapport « Chiffres clés ».

    build_docx(root, year, out_dir, mode="draft" | "final") -> (docx | None, summary)

Chaîne :
1. reconstruit le master canonique et sa certification (pipeline.run) ;
2. résout chaque valeur affichée sur les seules lignes PRODUCTION
   (même frontière que pipeline.publication : aucun repli VALIDATION,
   DIAGNOSTIC ou AUDIT, jamais de zéro par défaut) ;
3. produit le document :
   - avec le template Word s'il est présent (remplacement des {{TOKEN}}) ;
   - sinon, un document complet généré depuis referentiel/indicateurs.csv ;
4. écrit le préflight (une ligne par valeur, carte, graphique ou texte).

En mode final, tout élément requis manquant ou non publiable bloque la
génération : le préflight est écrit, le DOCX ne l'est pas.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
CONFIG_PATH = HERE / "reporting_config.json"
TOKEN_RE = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
NBSP = " "
EPCI = ["CACEM_EPCI", "CAESM_EPCI", "CAP_NORD_EPCI"]

PREFLIGHT_FIELDS = ["token", "indicator_id", "libelle", "territory", "perimeter_id", "page", "status",
                    "required", "blocking", "value", "source", "note"]
# Statuts du préflight qui interdisent la version finale.
BLOCKING = {"MISSING_VALUE", "NOT_PUBLISHABLE", "MAP_MISSING", "BLOCKING_ANOMALY", "UNKNOWN_TOKEN",
            "NOT_CERTIFIED"}
# Statut de couverture (pipeline.certification) -> statut du préflight.
COVERAGE_TO_STATUS = {
    "LACUNE_DECLAREE": "KNOWN_ABSENCE",
    "NON_APPLICABLE": "NOT_APPLICABLE",
    "REFERENCE_ONLY": "REFERENCE",
    "REFERENCE_EXTERNE": "REFERENCE",
}
DECIMALS = {"int_space": 0, "m3_int": 0, "km0": 0, "pct0": 0, "pct1": 1, "pct2": 2, "decimal1": 1,
            "decimal2": 2, "euro2": 2, "euro_m3": 2, "kg_abonne": 1}


def load_config(path: Path = CONFIG_PATH) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


# ------------------------------------------------------------------ formats
def format_number(value, fmt: str) -> str:
    """Format français centralisé : espace insécable des milliers, virgule décimale."""
    if fmt == "text":
        return str(value)
    v = float(value)
    d = DECIMALS.get(fmt, 0)
    s = f"{v:,.{d}f}"
    return s.replace(",", "\x00").replace(".", ",").replace("\x00", NBSP)


def format_for_unit(cfg: dict, unit: str) -> str:
    return cfg["formats_by_unit"].get(unit, cfg["default_format"])


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ------------------------------------------------------------------ résolution
@dataclass
class Resolution:
    indicator_id: str
    perimeter_id: str
    status: str          # OK, KNOWN_ABSENCE, NOT_APPLICABLE, REFERENCE, NOT_PUBLISHABLE, MISSING_VALUE, OPTIONAL_MISSING, NOT_EXPECTED
    text: str
    required: bool = False
    value: str = ""
    source: str = ""
    note: str = ""

    @property
    def ok(self) -> bool:
        return self.status == "OK"


class Context:
    """Données d'un millésime prêtes à publier."""

    def __init__(self, root: Path, year: int, cfg: dict, out: Path):
        if str(root) not in sys.path:
            sys.path.insert(0, str(root))
        from pipeline.master import Master
        from pipeline.referentiel import Referentiel, read_csv
        from pipeline.run import run

        self.root, self.year, self.cfg, self.out = root, year, cfg, out
        self.ref = Referentiel.load(root)
        res = run(year, root / "outputs" / str(year), root)
        self.pipeline_out: Path = res["out"]
        self.cert: dict = res["certification"]
        self.manifest: dict = res["manifest"]
        self.facts = read_csv(self.pipeline_out / f"fact_indicateur_master_{year}.csv")
        self.anomalies = read_csv(self.pipeline_out / f"anomalies_{year}.csv")
        self.cells = {(c["indicator_id"], c["perimeter_id"]): c
                      for c in read_csv(self.pipeline_out / f"couverture_{year}.csv")}
        self.prod, self.conflicts = production_index(self.facts)
        try:
            prev = Master(self.ref, year - 1).build().facts
        except Exception:  # millésime précédent sans source : pas de comparaison
            prev = []
        self.prev_prod, _ = production_index(prev)
        self.maps_dir = out / "assets_manual"

    # -- valeurs
    def unit(self, iid: str) -> str:
        return self.ref.indicateurs.get(iid, {}).get("unite", "")

    def libelle(self, iid: str) -> str:
        return self.ref.indicateurs.get(iid, {}).get("libelle", iid)

    def resolve(self, iid: str, pid: str) -> Resolution:
        texts = self.cfg["cell_texts"]
        cell = self.cells.get((iid, pid))
        required = bool(cell and cell["statut_attente"] == "EXPECTED")
        if (iid, pid) in self.conflicts:
            return Resolution(iid, pid, "NOT_PUBLISHABLE", texts["A_VALIDER"], required,
                              note="Plusieurs valeurs PRODUCTION divergentes : " + self.conflicts[(iid, pid)])
        row = self.prod.get((iid, pid))
        if row is not None and _num(row["value"]) is not None:
            return Resolution(iid, pid, "OK", format_number(row["value"], format_for_unit(self.cfg, self.unit(iid))),
                              required, row["value"], f"{row['source']} {row['source_file']}".strip(), row.get("note", ""))
        if cell is None:
            return Resolution(iid, pid, "NOT_EXPECTED", texts["NOT_EXPECTED"])
        cov = cell["statut_couverture"]
        note = cell.get("explication", "")
        if cov in COVERAGE_TO_STATUS:
            status = COVERAGE_TO_STATUS[cov]
            text = texts.get(cov, texts["MANQUANT"]) if status != "REFERENCE" else texts["NOT_EXPECTED"]
            return Resolution(iid, pid, status, text, required, note=note)
        if cov == "A_VALIDER":
            return Resolution(iid, pid, "NOT_PUBLISHABLE" if required else "OPTIONAL_MISSING",
                              texts["A_VALIDER"], required, note=note)
        return Resolution(iid, pid, "MISSING_VALUE" if required else "OPTIONAL_MISSING",
                          texts["MANQUANT"], required, note=note)

    def evolution(self, iid: str, pid: str = "MARTINIQUE") -> str:
        """Évolution N/N-1 sur deux valeurs PRODUCTION ; en points pour un pourcentage."""
        cur, prev = self.prod.get((iid, pid)), self.prev_prod.get((iid, pid))
        a, b = _num(cur["value"]) if cur else None, _num(prev["value"]) if prev else None
        if a is None or b is None:
            return ""
        if self.unit(iid).startswith("%"):
            return f"{'+' if a - b >= 0 else '−'}{format_number(abs(a - b), 'decimal1')} pt"
        if b == 0:
            return ""
        pct = (a - b) / b * 100
        return f"{'+' if pct >= 0 else '−'}{format_number(abs(pct), 'decimal1')} %"

    def dynamic_text(self, key: str) -> tuple[str, bool]:
        rule = self.cfg["dynamic_texts"][key]
        iid, pid = rule["indicator_id"], rule["perimeter_id"]
        cur, prev = self.prod.get((iid, pid)), self.prev_prod.get((iid, pid))
        a, b = _num(cur["value"]) if cur else None, _num(prev["value"]) if prev else None
        fmt = {"year": self.year, "prev_year": self.year - 1}
        if a is None or not b:
            return rule["unavailable"].format(**fmt), False
        pct = (a - b) / b * 100
        fmt.update(pct=format_number(abs(pct), "decimal1"),
                   signed_pct=("+" if pct >= 0 else "−") + format_number(abs(pct), "decimal1"))
        if abs(pct) < rule["stable_threshold_pct"]:
            return rule["stable"].format(**fmt), True
        return rule["up" if pct > 0 else "down"].format(**fmt), True

    # -- assets
    def map_path(self, token: str) -> Path | None:
        name = self.cfg["asset_filename_by_token"]["{{" + token + "}}"].format(year=self.year)
        p = self.maps_dir / name
        return p if p.exists() else None

    def manual_asset(self, token: str) -> Path | None:
        for ext in (".png", ".jpg", ".jpeg"):
            p = self.maps_dir / f"{token.lower()}{ext}"
            if p.exists():
                return p
        return None

    def token(self, iid: str, pid: str) -> str:
        """Token canonique d'une valeur (sans accolades) : EP_019_CACEM, EP_003_MARTINIQUE."""
        return f"{iid}_{self.ref.territoire(pid).upper()}"

    def parse_indicator_token(self, token: str) -> tuple[str, str | None] | None:
        ids = self.ref.indicateurs
        if token in ids:
            return token, "MARTINIQUE"
        for iid in sorted(ids, key=len, reverse=True):
            if token.startswith(iid + "_"):
                suffix = token[len(iid) + 1:]
                return iid, self.ref.perimetre(suffix) or self.ref.perimetre(suffix + "_EPCI")
        return None


def production_index(facts: list[dict]) -> tuple[dict, dict]:
    """Lignes PRODUCTION indexées par (indicateur, périmètre).

    Même frontière que pipeline.publication, mais une clé à valeurs divergentes
    est écartée (et signalée) au lieu d'interrompre tout le rapport : le master
    la déclare déjà comme anomalie bloquante CONFLIT_PRODUCTION.
    """
    index: dict[tuple[str, str], dict] = {}
    conflicts: dict[tuple[str, str], str] = {}
    for r in facts:
        if r.get("record_role") != "PRODUCTION":
            continue
        key = (r["indicator_id"], r["perimeter_id"])
        if key in index and _num(index[key]["value"]) != _num(r["value"]):
            conflicts[key] = f"{index[key]['value']} ({index[key]['source']}) / {r['value']} ({r['source']})"
        index.setdefault(key, r)
    for key in conflicts:
        index.pop(key, None)
    return index, conflicts


# ------------------------------------------------------------------ préflight
class Preflight:
    def __init__(self, ctx: Context):
        self.ctx = ctx
        self.rows: list[dict] = []
        self._seen: set[str] = set()

    def add(self, token: str, status: str, *, iid: str = "", pid: str = "", required: bool = False,
            value: str = "", source: str = "", note: str = "") -> None:
        if token in self._seen:
            return
        self._seen.add(token)
        self.rows.append({
            "token": token, "indicator_id": iid, "libelle": self.ctx.libelle(iid) if iid else "",
            "territory": self.ctx.ref.territoire(pid) if pid else "", "perimeter_id": pid, "page": "",
            "status": status, "required": "oui" if required else "non",
            "blocking": "oui" if status in BLOCKING else "non", "value": value, "source": source, "note": note,
        })

    def add_resolution(self, token: str, r: Resolution) -> None:
        self.add(token, r.status, iid=r.indicator_id, pid=r.perimeter_id, required=r.required,
                 value=r.value, source=r.source, note=r.note)

    def add_global_checks(self) -> None:
        for a in self.ctx.anomalies:
            if a.get("severite") == "BLOQUANT":
                self.add(f"ANOMALIE:{a['type']}:{a.get('indicator_id', '')}:{a.get('perimeter_id', '')}",
                         "BLOCKING_ANOMALY", iid=a.get("indicator_id", ""), pid=a.get("perimeter_id", ""),
                         note=a.get("message", ""))
        if self.ctx.cert.get("status") == "NOT_CERTIFIED" and not self.blocking():
            self.add("CERTIFICATION", "NOT_CERTIFIED", note="Le millésime n'est pas certifié par le pipeline.")

    def blocking(self) -> list[dict]:
        return [r for r in self.rows if r["blocking"] == "oui"]

    def write(self, path: Path) -> None:
        from pipeline.referentiel import write_csv
        write_csv(path, self.rows, PREFLIGHT_FIELDS)


# ------------------------------------------------------------------ graphiques
def make_chart(ctx: Context, spec: dict, path: Path) -> Path | None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    iid = spec["indicator_id"]
    pids = EPCI + (["MARTINIQUE"] if spec.get("with_martinique") else [])
    res = [ctx.resolve(iid, pid) for pid in pids]
    if not any(r.ok for r in res):
        return None
    labels = [ctx.ref.territoire(p).replace("_", " ") for p in pids]
    fig, ax = plt.subplots(figsize=(6.3, 2.8), dpi=200)
    for x, (pid, r) in enumerate(zip(pids, res)):
        if r.ok:
            v = float(r.value)
            ax.bar(x, v, width=0.6, color="#0b4f6c" if pid == "MARTINIQUE" else "#2e86ab")
            ax.text(x, v, r.text, ha="center", va="bottom", fontsize=8)
        else:
            ax.text(x, 0, r.text, ha="center", va="bottom", fontsize=8, color="#777777", style="italic")
    ax.set_xlim(-0.7, len(labels) - 0.3)
    ax.set_xticks(range(len(labels)), labels, fontsize=9)
    ax.set_title(spec["title"], fontsize=10, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.yaxis.grid(True, color="#e5e5e5", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.tick_params(axis="y", labelsize=8)
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path)
    plt.close(fig)
    return path


def build_charts(ctx: Context, pre: Preflight) -> dict[str, Path]:
    charts = {}
    for spec in ctx.cfg["charts"]:
        p = make_chart(ctx, spec, ctx.out / "charts" / f"{spec['id'].lower()}_{ctx.year}.png")
        if p:
            charts[spec["id"]] = p
            pre.add(spec["id"], "OK", iid=spec["indicator_id"], value=p.name)
        else:
            pre.add(spec["id"], "OPTIONAL_MISSING", iid=spec["indicator_id"],
                    note="Aucune valeur PRODUCTION : graphique non produit.")
    return charts


# ------------------------------------------------------------------ document généré
def _style_document(doc) -> None:
    from docx.shared import Pt, RGBColor
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x0B, 0x4F, 0x6C)


def _shade(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _write_cell(cell, r: Resolution, draft: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import RGBColor
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(r.text)
    if not r.ok:
        run.italic = True
        run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        if draft and r.status in BLOCKING:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _indicator_rows(ctx: Context, domain: str, perimeters: list[str]) -> list[str]:
    ids = []
    for iid, ind in ctx.ref.indicateurs.items():
        if ind["type"] != "PUBLICATION" or ind["domaine"] != domain:
            continue
        if any((iid, p) in ctx.cells or (iid, p) in ctx.prod for p in perimeters):
            ids.append(iid)
    return ids


def _other_perimeter_cells(ctx: Context, domain: str, table_perimeters: list[str]) -> list[tuple[str, str]]:
    keys = set()
    for (iid, pid), cell in ctx.cells.items():
        if pid not in table_perimeters and cell["statut_attente"] == "EXPECTED":
            keys.add((iid, pid))
    for (iid, pid) in ctx.prod:
        if pid not in table_perimeters and ctx.ref.perimetres.get(pid, {}).get("type") != "REFERENCE":
            keys.add((iid, pid))
    return sorted(k for k in keys if ctx.ref.indicateurs.get(k[0], {}).get("domaine") == domain
                  and ctx.ref.indicateurs[k[0]]["type"] == "PUBLICATION")


def _insert_asset(doc, token: str, label: str, path: Path | None, draft: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Cm
    if path:
        doc.add_picture(str(path), width=Cm(16))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(label)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    else:
        run = doc.add_paragraph().add_run(f"[{label} : image à fournir — {token}]")
        run.italic = True
        if draft:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def generate_document(ctx: Context, pre: Preflight, charts: dict[str, Path], draft: bool):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt

    cfg, year = ctx.cfg, ctx.year
    doc = Document()
    _style_document(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    if draft:
        hp = sec.header.paragraphs[0]
        hr = hp.add_run(f"BROUILLON — document de travail non publiable — généré le "
                        f"{datetime.now():%d/%m/%Y %H:%M}")
        hr.bold = True
        hr.font.size = Pt(8)

    # Couverture
    doc.add_paragraph(cfg["title"], style="Title")
    doc.add_paragraph(f"Données {year}").runs[0].font.size = Pt(16)
    doc.add_paragraph(cfg["publisher"])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Synthèse
    c = ctx.cert
    doc.add_heading("Synthèse du millésime", level=1)
    doc.add_paragraph(
        f"Statut de certification : {c.get('status', '?')}. "
        f"{c.get('cellules', {}).get('PRODUCTION', 0)} valeurs publiables sur {c.get('cellules_attendues', 0)} attendues "
        f"({format_number(c.get('taux_production_pct') or 0, 'decimal1')} %), {c.get('cellules', {}).get('LACUNE_DECLAREE', 0)} absences déclarées.")
    doc.add_paragraph("Seules les valeurs validées pour la production sont publiées. Une valeur absente "
                      "n'est jamais remplacée par zéro : elle est signalée comme « non produit » (absence "
                      "documentée par le service) ou « non disponible ».")

    tp = cfg["table_perimeters"]
    maps_cfg = cfg["maps"]
    for section in cfg["sections"]:
        domain = section["domain"]
        doc.add_heading(section["title"], level=1)
        for key in section.get("dynamic_texts", []):
            text, ok = ctx.dynamic_text(key)
            doc.add_paragraph(text)
            pre.add(key, "OK" if ok else "OPTIONAL_MISSING", note="" if ok else text)

        ids = _indicator_rows(ctx, domain, tp)
        if ids:
            headers = ["Indicateur", "Unité"] + [ctx.ref.territoire(p).replace("_", " ") for p in tp] + [f"Évol. {year - 1}"]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for i, h in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.text = h
                cell.paragraphs[0].runs[0].bold = True
                _shade(cell, "D9E7EE")
            for iid in ids:
                row = table.add_row().cells
                row[0].text = ctx.libelle(iid)
                row[1].text = ctx.unit(iid)
                for j, pid in enumerate(tp):
                    r = ctx.resolve(iid, pid)
                    _write_cell(row[2 + j], r, draft)
                    if r.status != "NOT_EXPECTED":
                        pre.add_resolution(ctx.token(iid, pid), r)
                row[-1].text = ctx.evolution(iid)
            for row in table.rows:
                row.cells[0].width = Cm(6.5)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8.5)

        others = _other_perimeter_cells(ctx, domain, tp)
        if others:
            doc.add_heading("Périmètres de service", level=2)
            t2 = doc.add_table(rows=1, cols=4)
            t2.style = "Table Grid"
            for i, h in enumerate(["Indicateur", "Périmètre", "Valeur", "Unité"]):
                t2.rows[0].cells[i].text = h
            for iid, pid in others:
                r = ctx.resolve(iid, pid)
                row = t2.add_row().cells
                row[0].text = ctx.libelle(iid)
                row[1].text = ctx.ref.territoire(pid).replace("_", " ")
                _write_cell(row[2], r, draft)
                row[3].text = ctx.unit(iid)
                pre.add_resolution(ctx.token(iid, pid), r)

        for spec in cfg["charts"]:
            iid = spec["indicator_id"]
            if ctx.ref.indicateurs.get(iid, {}).get("domaine") == domain and spec["id"] in charts:
                doc.add_picture(str(charts[spec["id"]]), width=Cm(15.5))

        for token in section.get("maps", []):
            m = maps_cfg[token]
            path = ctx.map_path(token)
            _insert_asset(doc, token, m["label"], path, draft)
            pre.add(token, "OK" if path else ("MAP_MISSING" if m.get("required") else "OPTIONAL_MISSING"),
                    required=m.get("required", False), value=path.name if path else "")

    # Absences et limites
    notes = [r for r in pre.rows if r["status"] in ("KNOWN_ABSENCE", "NOT_PUBLISHABLE") and r["indicator_id"]]
    if notes:
        doc.add_heading("Absences et limites", level=1)
        for r in notes:
            label = "Non produit" if r["status"] == "KNOWN_ABSENCE" else "À valider"
            doc.add_paragraph(f"{r['libelle']} — {r['territory'].replace('_', ' ')} : {label}. {r['note']}".strip(),
                              style="List Bullet")

    # Traçabilité
    doc.add_heading("Sources et traçabilité", level=1)
    doc.add_paragraph(f"Exécution {ctx.manifest.get('run_id', '')} du {ctx.manifest.get('finished_at', '')[:10]} — "
                      f"logiciel {_version(ctx.root)}.")
    for s in ctx.manifest.get("sources", []):
        doc.add_paragraph(f"{s['fichier']} (SHA-256 {s['sha256'][:12]}…)", style="List Bullet")
    return doc


# ------------------------------------------------------------------ template Word
def _iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc):
    yield from _iter_paragraphs(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from _iter_paragraphs(part)


def fill_template(ctx: Context, pre: Preflight, charts: dict[str, Path], template: Path, draft: bool):
    """Remplace les {{TOKEN}} du template ; une image remplace un paragraphe qui ne contient que son token."""
    from docx import Document
    from docx.shared import Cm

    doc = Document(str(template))

    def resolve(token: str) -> tuple[str, Path | None]:
        if token == "ANNEE":
            pre.add(token, "OK", value=str(ctx.year))
            return str(ctx.year), None
        if token == "ANNEE_PRECEDENTE":
            pre.add(token, "OK", value=str(ctx.year - 1))
            return str(ctx.year - 1), None
        if "{{" + token + "}}" in ctx.cfg["asset_filename_by_token"]:
            m = ctx.cfg["maps"].get(token, {})
            p = ctx.map_path(token)
            pre.add(token, "OK" if p else ("MAP_MISSING" if m.get("required") else "OPTIONAL_MISSING"),
                    required=m.get("required", False), value=p.name if p else "")
            return (f"[{m.get('label', token)} : image à fournir]" if not p else ""), p
        if token in charts:
            return "", charts[token]
        if token in ctx.cfg["dynamic_texts"]:
            text, ok = ctx.dynamic_text(token)
            pre.add(token, "OK" if ok else "OPTIONAL_MISSING", note="" if ok else text)
            return text, None
        parsed = ctx.parse_indicator_token(token)
        if parsed and parsed[1]:
            r = ctx.resolve(*parsed)
            pre.add_resolution(token, r)
            return r.text, None
        asset = ctx.manual_asset(token)
        if asset:
            pre.add(token, "OK", value=asset.name)
            return "", asset
        if token.startswith("CHART_"):
            return ctx.cfg["cell_texts"]["MANQUANT"], None
        pre.add(token, "UNKNOWN_TOKEN", note="Token non reconnu : ni indicateur, ni carte, ni graphique, "
                                             "ni image déposée dans assets_manual/.")
        return "{{" + token + "}}", None

    for p in _all_paragraphs(doc):
        if "{{" not in p.text or not p.runs:
            continue
        full = "".join(r.text for r in p.runs)
        tokens = TOKEN_RE.findall(full)
        if not tokens:
            continue
        first = p.runs[0]
        for r in p.runs[1:]:
            r.text = ""
        if len(tokens) == 1 and full.strip() == "{{" + tokens[0] + "}}":
            text, image = resolve(tokens[0])
            first.text = "" if image else text
            if image:
                first.add_picture(str(image), width=Cm(16))
            continue
        first.text = TOKEN_RE.sub(lambda m: resolve(m.group(1))[0], full)
    return doc


# ------------------------------------------------------------------ point d'entrée
def _version(root: Path) -> str:
    p = root / "VERSION"
    return p.read_text(encoding="utf-8").strip() if p.exists() else "?"


def build_docx(root: Path, year: int, out: Path, mode: str = "draft",
               config: dict | None = None) -> tuple[Path | None, dict]:
    if mode not in ("draft", "final"):
        raise ValueError(f"mode inconnu : {mode}")
    root, out = Path(root), Path(out)
    out.mkdir(parents=True, exist_ok=True)
    cfg = config or load_config()
    draft = mode == "draft"
    ctx = Context(root, year, cfg, out)
    pre = Preflight(ctx)
    charts = build_charts(ctx, pre)

    template = root / cfg["template_path"]
    if template.exists():
        doc = fill_template(ctx, pre, charts, template, draft)
    else:
        doc = generate_document(ctx, pre, charts, draft)
    pre.add_global_checks()
    pre.write(out / f"preflight_report_{year}.csv")

    quality = out / "quality"
    quality.mkdir(exist_ok=True)
    for name in (f"couverture_{year}.csv", f"anomalies_{year}.csv", f"certification_{year}.json",
                 f"statut_indicateurs_{year}.csv", "run_manifest.json"):
        src = ctx.pipeline_out / name
        if src.exists():
            shutil.copy2(src, quality / name)

    blocking = pre.blocking()
    counts: dict[str, int] = {}
    for r in pre.rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    summary = {
        "year": year, "mode": mode, "generated_at": datetime.now(timezone.utc).isoformat(),
        "template_used": str(template) if template.exists() else None,
        "certification": ctx.cert.get("status"), "taux_production_pct": ctx.cert.get("taux_production_pct"),
        "blocking_count": len(blocking), "status_counts": counts, "charts": len(charts),
        "blocking_examples": [f"{r['token']} : {r['status']}" for r in blocking[:10]],
    }
    docx_path = None
    if draft or not blocking:
        suffix = "_BROUILLON" if draft else ""
        docx_path = out / (cfg["output_name"].format(year=year) + suffix + ".docx")
        doc.save(str(docx_path))
    summary["docx"] = str(docx_path) if docx_path else None
    (out / f"report_summary_{year}.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2),
                                                     encoding="utf-8")
    return docx_path, summary
