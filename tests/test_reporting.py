from __future__ import annotations

import csv
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("report_builder", ROOT / "modules/reporting/report_builder.py")
RB = importlib.util.module_from_spec(SPEC)
sys.modules["report_builder"] = RB
assert SPEC.loader is not None
SPEC.loader.exec_module(RB)
NBSP = " "


def read_preflight(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


class FormatTests(unittest.TestCase):
    def test_formats_francais(self):
        self.assertEqual(f"7{NBSP}564{NBSP}328", RB.format_number("7564328", "int_space"))
        self.assertEqual("98,8", RB.format_number(98.7696, "pct1"))
        self.assertEqual(f"1{NBSP}234,50", RB.format_number(1234.5, "euro2"))
        self.assertEqual("305", RB.format_number("305.1", "km0"))

    def test_format_par_unite(self):
        cfg = RB.load_config()
        self.assertEqual("pct1", RB.format_for_unit(cfg, "%"))
        self.assertEqual("decimal2", RB.format_for_unit(cfg, "m³/km/j"))
        self.assertEqual(cfg["default_format"], RB.format_for_unit(cfg, "unité inconnue"))


class ProductionBoundaryTests(unittest.TestCase):
    def row(self, role, value, source="A"):
        return {"indicator_id": "EP_005", "perimeter_id": "CAESM_EPCI", "record_role": role,
                "value": value, "source": source}

    def test_seules_les_lignes_production_sont_publiables(self):
        index, conflicts = RB.production_index([self.row("DIAGNOSTIC", "10"), self.row("VALIDATION", "11")])
        self.assertEqual({}, index)
        self.assertEqual({}, conflicts)

    def test_valeurs_production_divergentes_ecartees(self):
        index, conflicts = RB.production_index([self.row("PRODUCTION", "10"), self.row("PRODUCTION", "12", "B")])
        self.assertNotIn(("EP_005", "CAESM_EPCI"), index)
        self.assertIn(("EP_005", "CAESM_EPCI"), conflicts)

    def test_doublon_identique_accepte(self):
        index, conflicts = RB.production_index([self.row("PRODUCTION", "10"), self.row("PRODUCTION", "10.0", "B")])
        self.assertIn(("EP_005", "CAESM_EPCI"), index)
        self.assertEqual({}, conflicts)


class BuildTests(unittest.TestCase):
    """Génération réelle sur le millésime 2024 (sources versionnées du dépôt)."""

    def test_brouillon_genere_et_preflight(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx, summary = RB.build_docx(ROOT, 2024, Path(tmp), mode="draft")
            self.assertTrue(docx and docx.exists())
            rows = read_preflight(Path(tmp) / "preflight_report_2024.csv")
            self.assertTrue(rows)
            self.assertEqual(summary["blocking_count"], sum(r["blocking"] == "oui" for r in rows))
            # Une absence déclarée n'est jamais bloquante ni convertie en zéro.
            absences = [r for r in rows if r["status"] == "KNOWN_ABSENCE"]
            self.assertTrue(absences)
            self.assertTrue(all(r["blocking"] == "non" and r["value"] == "" for r in absences))
            self.assertTrue((Path(tmp) / "quality" / "couverture_2024.csv").exists())

    def test_final_bloque_sans_document(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx, summary = RB.build_docx(ROOT, 2024, Path(tmp), mode="final")
            self.assertGreater(summary["blocking_count"], 0)
            self.assertIsNone(docx)
            self.assertEqual([], list(Path(tmp).glob("*.docx")))

    def test_template_remplace_les_tokens(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            tpl = Document()
            p = tpl.add_paragraph("Rapport ")
            p.add_run("{{AN")          # token coupé entre deux runs, comme dans Word
            p.add_run("NEE}}")
            tpl.add_paragraph("Abonnés CAESM : {{EP_005_CAESM}} ; CACEM : {{EP_006_CACEM}}")
            tpl.add_paragraph("{{MAP_CAPTAGES_AEP}}")
            tpl.add_paragraph("{{TOKEN_INCONNU}}")
            tpl_path = tmp / "template.docx"
            tpl.save(str(tpl_path))
            cfg = RB.load_config()
            cfg["template_path"] = str(tpl_path)
            docx, summary = RB.build_docx(ROOT, 2024, tmp / "out", mode="draft", config=cfg)
            text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
            self.assertIn("Rapport 2024", text)
            self.assertIn(f"Abonnés CAESM : 63{NBSP}269", text)
            self.assertIn("CACEM : non produit", text)
            self.assertNotIn("{{EP_005_CAESM}}", text)
            status = {r["token"]: r["status"] for r in read_preflight(tmp / "out" / "preflight_report_2024.csv")}
            self.assertEqual("OK", status["EP_005_CAESM"])
            self.assertEqual("KNOWN_ABSENCE", status["EP_006_CACEM"])
            self.assertEqual("MAP_MISSING", status["MAP_CAPTAGES_AEP"])
            self.assertEqual("UNKNOWN_TOKEN", status["TOKEN_INCONNU"])
            self.assertEqual(str(tpl_path), summary["template_used"])


class ConfigTests(unittest.TestCase):
    def test_graphiques_et_textes_referencent_des_indicateurs_connus(self):
        cfg = RB.load_config()
        with (ROOT / "referentiel/indicateurs.csv").open(encoding="utf-8-sig") as f:
            ids = {r["indicator_id"] for r in csv.DictReader(f, delimiter=";")}
        for spec in cfg["charts"]:
            self.assertIn(spec["indicator_id"], ids)
        for rule in cfg["dynamic_texts"].values():
            self.assertIn(rule["indicator_id"], ids)
        for section in cfg["sections"]:
            for token in section.get("maps", []):
                self.assertIn("{{" + token + "}}", cfg["asset_filename_by_token"])

    def test_config_json_valide(self):
        json.loads((ROOT / "modules/reporting/reporting_config.json").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
