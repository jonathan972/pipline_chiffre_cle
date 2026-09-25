#!/usr/bin/env python3
"""Initialisation (une seule fois) du référentiel canonique depuis le dictionnaire 2022.

Produit :
- referentiel/indicateurs.csv : les 80 indicateurs de publication + les variables support ;
- referentiel/attentes.csv    : ce qui doit exister pour chaque millésime (indicateur × périmètre).

Ces deux CSV deviennent ensuite la source de vérité et sont édités à la main
(Excel ou éditeur texte). Ce script n'est à relancer que pour repartir de zéro :
il refuse d'écraser un fichier existant sans --force.
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DICT = ROOT / "docs" / "Dictionnaire_indicateurs_ODE_Martinique_2022.xlsx"
OUT = ROOT / "referentiel"

DOMAINES = {
    "Ressource": "RESSOURCE",
    "Eau potable": "EAU_POTABLE",
    "Assainissement collectif": "ASSAINISSEMENT_COLLECTIF",
    "Assainissement non collectif": "ANC",
    "Tarifs": "TARIFS",
}

# Indicateurs dont la valeur Martinique est la somme des trois EPCI.
ADDITIFS = {
    "EP_005", "EP_006", "EP_007", "EP_014", "EP_015", "EP_017",
    "AC_001", "AC_002", "AC_012", "AC_014", "AC_015",
    "ANC_002", "ANC_003", "ANC_010", "ANC_011",
}

# Variables élémentaires utiles aux calculs ou aux contrôles, mais qui ne sont
# pas des indicateurs publiés isolément dans le rapport.
VARIABLES_SUPPORT = [
    ("POP_001", "DEMOGRAPHIE", "Population municipale INSEE", "hab", "SOMME"),
    ("VAR_EP_VOL_PRELEVE_LOCAL", "EAU_POTABLE", "Volume prélevé déclaré dans le RAD (contrôle de la BNPE)", "m³/an", "SOMME"),
    ("VAR_EP_VOL_MIS_DISTRIBUTION", "EAU_POTABLE", "Volume mis en distribution", "m³/an", "NON_ADDITIF"),
    ("VAR_EP_VOL_COMPTABILISE", "EAU_POTABLE", "Volume comptabilisé (hydraulique)", "m³/an", "NON_ADDITIF"),
    ("VAR_EP_VOL_FACTURE_CLIENTELE", "EAU_POTABLE", "Volume facturé selon le tableau clientèle du RAD", "m³/an", "NON_ADDITIF"),
    ("VAR_EP_VOL_FACTURE_HYDRAULIQUE", "EAU_POTABLE", "Volume facturé selon le bilan hydraulique du RAD", "m³/an", "NON_ADDITIF"),
    ("VAR_EP_VOL_VENDU_RPQS", "EAU_POTABLE", "Volume vendu déclaré dans le RPQS / SISPEA", "m³/an", "NON_ADDITIF"),
    ("VAR_AC_POP_DESSERVIE", "ASSAINISSEMENT_COLLECTIF", "Population desservie par l'assainissement collectif", "hab", "SOMME"),
    ("VAR_AC_VOL_ASSUJETTI", "ASSAINISSEMENT_COLLECTIF", "Volume assujetti à la redevance assainissement", "m³/an", "NON_ADDITIF"),
    ("VAR_AC_VOL_TRAITE", "ASSAINISSEMENT_COLLECTIF", "Volume traité en station", "m³/an", "NON_ADDITIF"),
    ("VAR_AC_NB_STEU_EXPLOITANT", "ASSAINISSEMENT_COLLECTIF", "Nombre de stations déclaré par l'exploitant (RAD)", "nb", "NON_ADDITIF"),
    ("VAR_AC_BOUES_PORTAIL", "ASSAINISSEMENT_COLLECTIF", "Boues déclarées au portail assainissement (contrôle de AC_012)", "tMS/an", "SOMME"),
    ("VAR_ANC_CTRL_CONCEPTION", "ANC", "Contrôles de conception ANC", "contrôles/an", "SOMME"),
    ("VAR_ANC_CTRL_EXECUTION", "ANC", "Contrôles de bonne exécution ANC", "contrôles/an", "SOMME"),
    ("VAR_ANC_DIAG_VENTE", "ANC", "Diagnostics ANC réalisés lors de ventes", "contrôles/an", "SOMME"),
]

# Libellés de territoire du rapport 2022 -> périmètre canonique.
TERRITOIRES = {
    "Martinique": "MARTINIQUE",
    "Martinique (texte)": "MARTINIQUE",
    "CACEM": "CACEM_EPCI",
    "CAESM": "CAESM_EPCI",
    "CAP Nord": "CAP_NORD_EPCI",
    "France": "FRANCE",
    "Guadeloupe": "GUADELOUPE",
}

# Indicateurs de performance publiés par contrat/service, pas par EPCI.
PAR_CONTRAT = {"EP_019", "EP_020", "EP_021"}


def attentes_contrat(indicator_id: str, label: str) -> list[tuple[str, str, str, str]]:
    """(perimeter_id, annee_debut, annee_fin, note) pour les indicateurs par contrat."""
    if label.startswith("CAESM"):
        return [("EX_SICSM_DSP_2015_2027", "2022", "2027",
                 "Publié « CAESM (+ Robert/Trinité) » : non ventilable entre CAESM et Robert-Trinité.")]
    if label == "CAP Nord":
        return [
            ("CAP_NORD_DSP_2020_2024", "2022", "2024", "Contrat SME CAP Nord (16 communes + secteurs Robert CN / Trinité CN)."),
            ("CAP_NORD_SAUR_2025_2027", "2025", "", "Contrat SAUR à partir de 2025 : à confirmer lors de la production 2025."),
        ]
    if label == "CACEM":
        return [("CACEM_EPCI", "2022", "", "Régie ODYSSI.")]
    return []


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    OUT.mkdir(exist_ok=True)
    targets = [OUT / "indicateurs.csv", OUT / "attentes.csv"]
    if any(p.exists() for p in targets) and not args.force:
        raise SystemExit("Le référentiel existe déjà : éditez les CSV, ou relancez avec --force.")

    wb = openpyxl.load_workbook(DICT, read_only=True)

    indicateurs = []
    acquisition = {}
    for r in wb["Catalogue"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        iid = r[0]
        acquisition[iid] = r[14] or ""
        indicateurs.append({
            "indicator_id": iid,
            "type": "PUBLICATION",
            "domaine": DOMAINES[r[1]],
            "libelle": r[3],
            "unite": r[5],
            "granularite": r[6],
            "agregation_martinique": "SOMME_EPCI" if iid in ADDITIFS else "NON_ADDITIF",
            "acquisition": acquisition[iid],
            "source_cible": r[10] or "",
            "code_officiel": "" if r[11] in (None, "—") else r[11],
        })
    for iid, dom, lib, unit, agg in VARIABLES_SUPPORT:
        indicateurs.append({
            "indicator_id": iid, "type": "SUPPORT", "domaine": dom, "libelle": lib, "unite": unit,
            "granularite": "", "agregation_martinique": "SOMME_EPCI" if agg == "SOMME" else "NON_ADDITIF",
            "acquisition": "", "source_cible": "", "code_officiel": "",
        })

    # Attentes : ce que le rapport 2022 publie, au grain du périmètre.
    cells: dict[tuple[str, str, str, str], str] = {}
    for r in wb["Validation_2022"].iter_rows(min_row=2, values_only=True):
        iid, _year, label = r[0], r[1], str(r[2] or "")
        if not iid:
            continue
        if iid in PAR_CONTRAT and label not in ("France", "Martinique (texte)"):
            for pid, deb, fin, note in attentes_contrat(iid, label):
                cells[(iid, pid, deb, fin)] = note
            continue
        pid = TERRITOIRES.get(label)
        if pid is None:
            # Déclinaisons (matériau, usine, commune, top 4…) : l'attente porte sur Martinique.
            pid, note = "MARTINIQUE", f"Déclinaison publiée : {label}."
        else:
            note = ""
        cells.setdefault((iid, pid, "2022", ""), note)

    attentes = []
    seen = {k[0] for k in cells}
    for ind in indicateurs:
        if ind["type"] != "PUBLICATION" or ind["indicator_id"] in seen:
            continue
        cells[(ind["indicator_id"], "MARTINIQUE", "2022", "")] = "Non chiffré dans la validation 2022 : attente par défaut."
    for (iid, pid, deb, fin), note in sorted(cells.items()):
        if pid in ("FRANCE", "GUADELOUPE"):
            statut = "REFERENCE_EXTERNE"
        elif "RÉFÉRENCE" in acquisition.get(iid, ""):
            statut = "REFERENCE_ONLY"
        else:
            statut = "EXPECTED"
        attentes.append({"indicator_id": iid, "perimeter_id": pid, "statut_attente": statut,
                         "annee_debut": deb, "annee_fin": fin, "note": note})

    for path, rows in ((targets[0], indicateurs), (targets[1], attentes)):
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=";")
            w.writeheader()
            w.writerows(rows)
        print(f"{path.relative_to(ROOT)} : {len(rows)} lignes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
